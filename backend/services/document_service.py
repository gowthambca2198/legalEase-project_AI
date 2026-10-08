import base64
from io import BytesIO
import os
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

from fpdf import FPDF

from backend.schemas import ExportRequest
from backend.utils.text_utils import (
    sanitize_text,
    terms_to_list,
)


class LegalDocumentService:
    """
    Responsible for TXT, DOCX and PDF generation.
    """

    # -------------------------------------------------
    # TXT
    # -------------------------------------------------

    def txt(
        self,
        request: ExportRequest,
    ) -> bytes:
        content = sanitize_text(
            request.content
        )

        header = (
            f"{request.document_type}\n"
            f"Effective Date: "
            f"{request.effective_date}\n\n"
        )

        result = (
            header
            + content
            + "\n"
        )

        return result.encode(
            "utf-8"
        )

    # -------------------------------------------------
    # DOCX
    # -------------------------------------------------

    def docx(
        self,
        request: ExportRequest,
    ) -> bytes:
        document = Document()
        font_name = request.branding_font

        section = document.sections[0]

        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

        normal_style = document.styles[
            "Normal"
        ]

        normal_style.font.name = (
            font_name
        )

        normal_style.font.size = Pt(11)

        if request.logo_base64:
            logo_bytes = base64.b64decode(
                request.logo_base64,
                validate=True,
            )
            logo_paragraph = document.add_paragraph()
            logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            logo_paragraph.add_run().add_picture(
                BytesIO(logo_bytes),
                width=Inches(1.1),
            )

        # Title
        title = document.add_paragraph()

        title.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        title_run = title.add_run(
            request.document_type.upper()
        )

        title_run.bold = True
        title_run.font.name = (
            font_name
        )
        title_run.font.size = Pt(16)

        # Branding
        branding = document.add_paragraph()

        branding.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        branding_run = branding.add_run(
            request.branding_name
        )

        branding_run.italic = True
        branding_run.font.size = Pt(9)

        # Metadata table
        metadata_table = document.add_table(
            rows=2,
            cols=2,
        )

        metadata_table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        metadata_table.style = "Table Grid"

        metadata_table.cell(
            0,
            0,
        ).text = "Effective Date"

        metadata_table.cell(
            0,
            1,
        ).text = request.effective_date

        metadata_table.cell(
            1,
            0,
        ).text = "Jurisdiction"

        metadata_table.cell(
            1,
            1,
        ).text = request.jurisdiction

        document.add_paragraph()

        # Main document
        clean_content = sanitize_text(
            request.content
        )

        blocks = re.split(
            r"\n\s*\n",
            clean_content,
        )

        for block in blocks:
            block = block.strip()

            if not block:
                continue

            lines = block.splitlines()

            first_line = lines[0].strip()

            heading_match = re.match(
                r"^(?:\d+[\.)]|\d+\s+-|[A-Z][A-Za-z ]+:)",
                first_line,
            )

            if heading_match:
                paragraph = document.add_paragraph()

                run = paragraph.add_run(
                    first_line
                )

                run.bold = True

                for line in lines[1:]:
                    if line.strip():
                        document.add_paragraph(
                            line.strip()
                        )

            else:
                paragraph = document.add_paragraph(
                    block
                )

                paragraph.paragraph_format.space_after = Pt(
                    7
                )

        # Terms table
        terms = terms_to_list(
            request.terms
        )

        if terms:
            document.add_heading(
                "Requested Terms",
                level=2,
            )

            table = document.add_table(
                rows=1,
                cols=2,
            )

            table.style = "Table Grid"

            table.alignment = (
                WD_TABLE_ALIGNMENT.CENTER
            )

            table.rows[0].cells[0].text = "No."
            table.rows[0].cells[1].text = "Term"

            for index, term in enumerate(
                terms,
                start=1,
            ):
                cells = table.add_row().cells

                cells[0].text = str(index)
                cells[1].text = term

        # Footer
        footer = section.footer.paragraphs[0]

        footer.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        footer.add_run(
            f"{request.branding_name} | "
            "AI-assisted draft | "
            "Review before use"
        )

        output = BytesIO()

        document.save(output)

        return output.getvalue()

    # -------------------------------------------------
    # PDF
    # -------------------------------------------------

    def pdf(
        self,
        request: ExportRequest,
    ) -> bytes:
        content = sanitize_text(
            request.content
        )

        class LegalPDF(FPDF):

            def header(self):
                if self.logo_bytes:
                    self.image(
                        BytesIO(self.logo_bytes),
                        x=(self.w - 30) / 2,
                        y=4,
                        w=30,
                        h=12,
                        keep_aspect_ratio=True,
                    )
                    brand_y = 19
                else:
                    brand_y = 7

                self.set_font(
                    self.branding_font,
                    "B",
                    9,
                )

                self.set_xy(self.l_margin, brand_y)
                self.cell(
                    self.epw,
                    7,
                    self.branding,
                    align="C",
                )

            def footer(self):
                self.set_y(-15)

                self.set_font(
                    self.branding_font,
                    "",
                    8,
                )

                self.cell(
                    0,
                    8,
                    (
                        f"{self.branding} | "
                        "AI-assisted draft | "
                        f"Page {self.page_no()}"
                    ),
                    align="C",
                )

        pdf = LegalPDF()

        def add_full_width_text(
            text: str,
            line_height: float,
            **options,
        ) -> None:
            # fpdf2 advances the cursor to the right edge after multi_cell.
            # Reset it before the next full-width block to avoid a zero-width
            # line on subsequent calls.
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(
                pdf.epw,
                line_height,
                text,
                **options,
            )

        pdf.branding = (
            request.branding_name
            or "LegalEase"
        )
        core_fonts = {
            "Times New Roman": "Times",
            "Arial": "Helvetica",
            "Courier New": "Courier",
        }[request.branding_font]
        windows_fonts = Path(
            os.environ.get("WINDIR", "C:/Windows")
        ) / "Fonts"
        font_files = {
            "Times New Roman": ("times.ttf", "timesbd.ttf"),
            "Arial": ("arial.ttf", "arialbd.ttf"),
            "Courier New": ("cour.ttf", "courbd.ttf"),
        }
        regular_font, bold_font = (
            windows_fonts / name
            for name in font_files[request.branding_font]
        )

        if regular_font.is_file() and bold_font.is_file():
            pdf.add_font(
                "LegalEaseUnicode",
                "",
                str(regular_font),
            )
            pdf.add_font(
                "LegalEaseUnicode",
                "B",
                str(bold_font),
            )
            pdf.branding_font = "LegalEaseUnicode"
        else:
            pdf.branding_font = core_fonts
        pdf.logo_bytes = (
            base64.b64decode(
                request.logo_base64,
                validate=True,
            )
            if request.logo_base64
            else b""
        )
        pdf.set_margins(
            10,
            32 if pdf.logo_bytes else 20,
            10,
        )

        pdf.set_auto_page_break(
            auto=True,
            margin=20,
        )

        pdf.add_page()

        # Title
        pdf.set_font(
            pdf.branding_font,
            "B",
            16,
        )

        add_full_width_text(
            request.document_type.upper(),
            9,
            align="C",
        )

        # Metadata
        pdf.set_font(
            pdf.branding_font,
            "",
            9,
        )

        add_full_width_text(
            (
                f"Effective Date: "
                f"{request.effective_date} | "
                f"Jurisdiction: "
                f"{request.jurisdiction}"
            ),
            6,
            align="C",
        )

        pdf.ln(5)

        # Main content
        blocks = re.split(
            r"\n\s*\n",
            content,
        )

        for block in blocks:
            block = block.strip()

            if not block:
                continue

            lines = block.splitlines()

            first_line = lines[0].strip()

            heading_match = re.match(
                r"^(?:\d+[\.)]|\d+\s+-|[A-Z][A-Za-z ]+:)",
                first_line,
            )

            if heading_match:
                pdf.set_font(
                    pdf.branding_font,
                    "B",
                    11,
                )

                add_full_width_text(
                    first_line,
                    7,
                )

                remainder = "\n".join(
                    lines[1:]
                ).strip()

                if remainder:
                    pdf.set_font(
                        pdf.branding_font,
                        "",
                        10,
                    )

                    add_full_width_text(
                        remainder,
                        6,
                    )

            else:
                pdf.set_font(
                    pdf.branding_font,
                    "",
                    10,
                )

                add_full_width_text(
                    block,
                    6,
                )

            pdf.ln(2)

        # Terms
        terms = terms_to_list(
            request.terms
        )

        if terms:
            pdf.set_font(
                pdf.branding_font,
                "B",
                12,
            )

            add_full_width_text(
                "Requested Terms",
                8,
            )

            pdf.set_font(
                pdf.branding_font,
                "",
                10,
            )

            for index, term in enumerate(
                terms,
                start=1,
            ):
                add_full_width_text(
                    f"{index}. {term}",
                    6,
                )

        output = pdf.output()

        return bytes(output)
